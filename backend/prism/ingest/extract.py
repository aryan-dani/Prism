"""Shared plain-text extraction for PDF / DOCX / HTML / Markdown / text.

Used by session uploads, permanent-KB upsert, and warranty-fixture ingest so
the parser story is one abstraction: HTML today, official PDF/DOCX when Kohler
hands us the corpus.
"""

from __future__ import annotations

from pathlib import Path


def extract_plain_text(path: Path, suffix: str | None = None) -> str:
    suffix = (suffix or path.suffix).lower()
    if suffix == ".pdf":
        from prism.ingest.parse_pdf import extract_text as extract_pdf_text

        return extract_pdf_text(path)
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    if suffix in {".html", ".htm"}:
        from bs4 import BeautifulSoup

        raw = path.read_text(encoding="utf-8", errors="replace")
        soup = BeautifulSoup(raw, "lxml")
        for tag in soup(["script", "style", "nav", "footer"]):
            tag.decompose()
        return soup.get_text("\n", strip=True)
    return path.read_text(encoding="utf-8", errors="replace")
