"""Excel (.xlsx) renderer -- only meaningful when the answer is genuinely
tabular (an explicit table) or has enough discrete facts/steps to be worth a
spreadsheet. `can_render` gates this so the UI only offers Excel when it
will actually produce something useful."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from prism.core.answer import CanonicalAnswer


def can_render(answer: CanonicalAnswer) -> bool:
    return answer.table is not None or len(answer.key_facts) >= 3 or len(answer.steps) >= 3


def render(answer: CanonicalAnswer) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Prism Answer"

    bold = Font(bold=True)

    ws.append(["Query", answer.query])
    ws.append(["Domain", answer.domain])
    ws.append(["Confidence", answer.confidence])
    ws.append(["Direct Answer", answer.direct_answer])
    ws.append([])

    if answer.table:
        header_row = ws.max_row + 1
        ws.append(answer.table.columns)
        for cell in ws[header_row]:
            cell.font = bold
        for row in answer.table.rows:
            ws.append(row)
    elif answer.key_facts:
        header_row = ws.max_row + 1
        ws.append(["Label", "Value", "Unit"])
        for cell in ws[header_row]:
            cell.font = bold
        for kf in answer.key_facts:
            ws.append([kf.label, kf.value, kf.unit or ""])
    elif answer.steps:
        header_row = ws.max_row + 1
        ws.append(["Step #", "Action"])
        for cell in ws[header_row]:
            cell.font = bold
        for i, step in enumerate(answer.steps, 1):
            ws.append([i, step])

    if answer.caveats:
        ws.append([])
        ws.append(["Caveats"])
        for c in answer.caveats:
            ws.append([c])

    if answer.sources:
        ws.append([])
        ws.append(["Sources"])
        for s in answer.sources:
            ws.append([s])

    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 12), 60)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
